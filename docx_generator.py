import os
from datetime import datetime
import docx
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.oxml import parse_xml
from docx.oxml.ns import nsdecls

# Color institucional FBioyF: #25A09C -> RGB(37, 160, 156)
COLOR_HEX = "25A09C"
COLOR_RGB = RGBColor(37, 160, 156)

def set_cell_background(cell, hex_color):
    """Establece el color de fondo de una celda de tabla."""
    tcPr = cell._element.get_or_add_tcPr()
    shd = parse_xml(f'<w:shd {nsdecls("w")} w:fill="{hex_color}"/>')
    tcPr.append(shd)

def agregar_encabezado_institucional(doc, logo_path="logo_fbioyf.png"):
    """Agrega la cabecera institucional con logo (si existe) y títulos en color institucional."""
    if os.path.exists(logo_path):
        # Creamos 3 columnas: Columna 0 (Logo), Columna 1 (Espaciador), Columna 2 (Texto)
        header_table = doc.add_table(rows=1, cols=3)
        header_table.alignment = WD_TABLE_ALIGNMENT.CENTER
        header_table.autofit = False
        
        cell_logo = header_table.rows[0].cells[0]
        cell_logo.width = Inches(1.3)
        p_logo = cell_logo.paragraphs[0]
        p_logo.alignment = WD_ALIGN_PARAGRAPH.LEFT
        p_logo.add_run().add_picture(logo_path, width=Inches(1.1))
        
        # Columna del medio (espaciador de 0.3 pulgadas)
        cell_space = header_table.rows[0].cells[1]
        cell_space.width = Inches(0.3)
        
        cell_text = header_table.rows[0].cells[2]
        cell_text.width = Inches(4.9)
        p_text = cell_text.paragraphs[0]
        p_text.alignment = WD_ALIGN_PARAGRAPH.LEFT
    else:
        p_text = doc.add_paragraph()
        p_text.alignment = WD_ALIGN_PARAGRAPH.CENTER

    run_title1 = p_text.add_run("UNIVERSIDAD NACIONAL DE ROSARIO\n")
    run_title1.font.name = "Arial"
    run_title1.font.size = Pt(11)
    run_title1.font.bold = True
    run_title1.font.color.rgb = COLOR_RGB

    run_title2 = p_text.add_run("Facultad de Ciencias Bioquímicas y Farmacéuticas\n")
    run_title2.font.name = "Arial"
    run_title2.font.size = Pt(12)
    run_title2.font.bold = True
    run_title2.font.color.rgb = COLOR_RGB

    run_sub = p_text.add_run("SISTEMA PANDA - PROGRAMA ANALÍTICO (PUPA)")
    run_sub.font.name = "Arial"
    run_sub.font.size = Pt(10)
    run_sub.font.italic = True

    doc.add_paragraph().paragraph_format.space_after = Pt(12)

def agregar_titulo_seccion(doc, texto):
    """Agrega un encabezado de sección formateado con el color institucional."""
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(14)
    p.paragraph_format.space_after = Pt(6)
    p.paragraph_format.keep_with_next = True
    
    run = p.add_run(texto)
    run.font.name = "Arial"
    run.font.size = Pt(12)
    run.font.bold = True
    run.font.color.rgb = COLOR_RGB

def agregar_campo(doc, etiqueta, valor):
    """Agrega un par etiqueta - valor formateado."""
    p = doc.add_paragraph()
    p.paragraph_format.space_after = Pt(4)
    
    run_lbl = p.add_run(f"{etiqueta}: ")
    run_lbl.font.name = "Arial"
    run_lbl.font.size = Pt(10)
    run_lbl.font.bold = True
    
    val_text = str(valor) if valor is not None and str(valor).strip() != "" else "No especificado"
    run_val = p.add_run(val_text)
    run_val.font.name = "Arial"
    run_val.font.size = Pt(10)

def generar_docx_programa(datos_materia, datos_formulario, ruta_salida="programa_analitico.docx", logo_path="logo_fbioyf.png"):
    """Genera el archivo Word completo combinando datos de BD y del formulario PUPA."""
    doc = docx.Document()

    # Márgenes de la página
    for section in doc.sections:
        section.top_margin = Inches(0.8)
        section.bottom_margin = Inches(0.8)
        section.left_margin = Inches(0.8)
        section.right_margin = Inches(0.8)

    # Header con logo e institución
    agregar_encabezado_institucional(doc, logo_path)

    # 1. INFORMACIÓN GENERAL (BD)
    agregar_titulo_seccion(doc, "1. INFORMACIÓN GENERAL DE LA ASIGNATURA")
    agregar_campo(doc, "Asignatura", datos_materia.get("nombre", ""))
    agregar_campo(doc, "Carrera/s", datos_materia.get("carrera", ""))
    agregar_campo(doc, "Código/s", datos_materia.get("codigo", ""))
    agregar_campo(doc, "Departamento / Área", datos_materia.get("departamento", ""))
    agregar_campo(doc, "Dictado", datos_materia.get("plan", ""))
    agregar_campo(doc, "Ubicación en el Ciclo", datos_materia.get("ciclo", ""))
    agregar_campo(doc, "Contenidos Básicos", datos_materia.get("contenidos_basicos", ""))
    agregar_campo(doc, "Alcances del Plan", datos_materia.get("alcances_plan", ""))
    agregar_campo(doc, "Docente Responsable", datos_materia.get("equipo_docente", ""))

    # 2. FUNDAMENTACIÓN
    agregar_titulo_seccion(doc, "2. FUNDAMENTACIÓN")
    agregar_campo(doc, "Fundamentación", datos_formulario.get("fundamentos", ""))

# 3. CARGA HORARIA
    agregar_titulo_seccion(doc, "3. CARGA HORARIA")
    
    # Aclaración de tipos de prácticas antes de la tabla
    p_info = doc.add_paragraph()
    p_info.paragraph_format.space_after = Pt(6)
    
    run_info_title = p_info.add_run("Tipos de Prácticas:\n")
    run_info_title.font.name = "Arial"
    run_info_title.font.size = Pt(9.5)
    run_info_title.font.bold = True
    
    p_info.add_run("• P1 (Prácticas de Ejercitación): Se incluyen aquí las prácticas de ejercitación: resolución de problemas a través de la reiteración de actividades (reconocimiento, exploración, ejercicios o problemas) tendientes a la incorporación de conceptos.\n")
    p_info.add_run("• P2 (Prácticas Experimentales): Se incluyen aquí las prácticas experimentales: actividades que promueven el desarrollo de habilidades propias basadas en la observación, exploración, experimentación y organización. Incluye prácticas de laboratorio, gabinete, en territorio o áulicas.\n")
    p_info.add_run("• P3 (Prácticas Integradoras Relacionadas con el Ejercicio Profesional): Se incluyen aquí las prácticas integradoras relacionadas con el ejercicio profesional: actividades que impliquen resolución de casos clínicos, simulación, elaboración de monografías, escritura de trabajos de divulgación, entre otras. En esta categoría también estarían incluidas análisis y discusión de trabajos científicos, búsqueda y análisis crítico de información en ciencias, entre otras.")
    
    for r in p_info.runs[1:]:
        r.font.name = "Arial"
        r.font.size = Pt(9)

    table_h = doc.add_table(rows=2, cols=5)
    table_h.alignment = WD_TABLE_ALIGNMENT.CENTER
    headers_h = ["Teóricas", "Prácticas P1", "Prácticas P2", "Prácticas P3", "Otras Actividades"]
    
    hdr_cells = table_h.rows[0].cells
    for i, title in enumerate(headers_h):
        hdr_cells[i].text = title
        set_cell_background(hdr_cells[i], COLOR_HEX)
        p = hdr_cells[i].paragraphs[0]
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        for run in p.runs:
            run.font.name = "Arial"
            run.font.size = Pt(9)
            run.font.bold = True
            run.font.color.rgb = RGBColor(255, 255, 255)

    vals_h = [
        f"{datos_formulario.get('num_h_teoricas', 0)} hs",
        f"{datos_formulario.get('num_h_p1', 0)} hs",
        f"{datos_formulario.get('num_h_p2', 0)} hs",
        f"{datos_formulario.get('num_h_p3', 0)} hs",
        f"{datos_formulario.get('num_h_otras', 0)} hs"
    ]
    val_cells = table_h.rows[1].cells
    for i, val in enumerate(vals_h):
        val_cells[i].text = val
        p = val_cells[i].paragraphs[0]
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        for run in p.runs:
            run.font.name = "Arial"
            run.font.size = Pt(9)

    doc.add_paragraph().paragraph_format.space_after = Pt(4)
    agregar_campo(doc, "Carga horaria total de la asignatura", f"{datos_formulario.get('total_horas', 0)} hs")
    agregar_campo(doc, "Detalle de Otras Actividades", datos_formulario.get("text_detalle_otros", ""))
    agregar_campo(doc, "Modalidad de Cursado", datos_formulario.get("radio_modalidad_virtual", ""))
    agregar_campo(doc, "Porcentaje Virtual (si es Mixto)", datos_formulario.get("input_porcentaje_mixto", ""))
    agregar_campo(doc, "Espacios de Consulta", datos_formulario.get("text_espacios_consulta", ""))

    # 4. PROGRAMA ANALÍTICO
    agregar_titulo_seccion(doc, "4. PROGRAMA ANALÍTICO Y CONTENIDOS")
    
    # Recibimos la lista estructurada de unidades en lugar de texto plano
    unidades_lista = datos_formulario.get("unidades_lista", [])
    
    if unidades_lista:
        for u in unidades_lista:
            # Subtítulo de la Unidad (Negrita y un poco más grande)
            p_uni = doc.add_paragraph()
            p_uni.paragraph_format.space_before = Pt(8)
            p_uni.paragraph_format.space_after = Pt(2)
            p_uni.paragraph_format.keep_with_next = True
            run_titulo_uni = p_uni.add_run(f"UNIDAD {u.get('numero', '')}: {u.get('nombre', '')}")
            run_titulo_uni.font.name = "Arial"
            run_titulo_uni.font.size = Pt(11)
            run_titulo_uni.font.bold = True
            run_titulo_uni.font.color.rgb = COLOR_RGB

            # Contenidos conceptuales en párrafo separado
            p_con = doc.add_paragraph()
            p_con.paragraph_format.space_after = Pt(2)
            run_lbl_con = p_con.add_run("Contenidos: ")
            run_lbl_con.font.bold = True
            run_lbl_con.font.name = "Arial"
            run_lbl_con.font.size = Pt(10)
            run_val_con = p_con.add_run(u.get('conceptos', 'No especificado'))
            run_val_con.font.name = "Arial"
            run_val_con.font.size = Pt(10)

            # Trabajos Prácticos en párrafo separado
            p_tp = doc.add_paragraph()
            p_tp.paragraph_format.space_after = Pt(2)
            run_lbl_tp = p_tp.add_run("Trabajos Prácticos / Seminarios: ")
            run_lbl_tp.font.bold = True
            run_lbl_tp.font.name = "Arial"
            run_lbl_tp.font.size = Pt(10)
            run_val_tp = p_tp.add_run(u.get('trabajos_practicos', 'No especificado'))
            run_val_tp.font.name = "Arial"
            run_val_tp.font.size = Pt(10)

            # Tipo de Prácticas en párrafo separado
            p_tipos = doc.add_paragraph()
            p_tipos.paragraph_format.space_after = Pt(8)
            run_lbl_tipos = p_tipos.add_run("Tipo de Prácticas: ")
            run_lbl_tipos.font.bold = True
            run_lbl_tipos.font.name = "Arial"
            run_lbl_tipos.font.size = Pt(10)
            run_val_tipos = p_tipos.add_run(u.get('tipos', 'No especificado'))
            run_val_tipos.font.name = "Arial"
            run_val_tipos.font.size = Pt(10)
    else:
        agregar_campo(doc, "Unidades Temáticas", "Sin unidades registradas")

    # 5. METODOLOGÍA Y ESTRATEGIAS DIDÁCTICAS
    agregar_titulo_seccion(doc, "5. METODOLOGÍA Y ESTRATEGIAS DIDÁCTICAS")
    agregar_campo(doc, "5.1. Teóricas", datos_formulario.get("metodologia_teoricas", ""))
    agregar_campo(doc, "5.2. Prácticas P1", datos_formulario.get("metodologia_p1", ""))
    agregar_campo(doc, "5.3. Prácticas P2", datos_formulario.get("metodologia_p2", ""))
    agregar_campo(doc, "5.4. Prácticas P3", datos_formulario.get("metodologia_p3", ""))

    # 6. CRITERIOS DE EVALUACIÓN
    agregar_titulo_seccion(doc, "6. CRITERIOS DE EVALUACIÓN")
    criterios = datos_formulario.get("criterios_evaluacion", [])
    if criterios:
        table_crit = doc.add_table(rows=1, cols=5)
        table_crit.alignment = WD_TABLE_ALIGNMENT.CENTER
        headers_crit = ["Criterio Evaluado", "Cursado", "Final Reg.", "Final Lib.", "N/A"]
        
        c_hdr = table_crit.rows[0].cells
        for i, h in enumerate(headers_crit):
            c_hdr[i].text = h
            set_cell_background(c_hdr[i], COLOR_HEX)
            p = c_hdr[i].paragraphs[0]
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            for run in p.runs:
                run.font.name = "Arial"
                run.font.size = Pt(8.5)
                run.font.bold = True
                run.font.color.rgb = RGBColor(255, 255, 255)

        for item in criterios:
            row_cells = table_crit.add_row().cells
            row_cells[0].text = item.get("nombre", "")
            row_cells[1].text = "X" if item.get("cursado") else ""
            row_cells[2].text = "X" if item.get("regulares") else ""
            row_cells[3].text = "X" if item.get("libres") else ""
            row_cells[4].text = "X" if item.get("no_aplica") else ""

            for i in range(5):
                p = row_cells[i].paragraphs[0]
                if i > 0:
                    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
                for run in p.runs:
                    run.font.name = "Arial"
                    run.font.size = Pt(8)

    doc.add_paragraph().paragraph_format.space_after = Pt(4)
    agregar_campo(doc, "Otros Criterios Específicos", datos_formulario.get("otros_criterios_eval", ""))

    # 7. CONDICIONES DE REGULARIDAD Y SISTEMA DE PROMOCIÓN
    agregar_titulo_seccion(doc, "7. CONDICIONES DE REGULARIDAD Y SISTEMA DE PROMOCIÓN")
    agregar_campo(doc, "7.1. Definición de las Actividades Obligatorias para acceder a la Regularidad", datos_formulario.get("actividades_obligatorias", ""))
    agregar_campo(doc, "7.2. % de Asistencia Exigido para Regularizar", datos_formulario.get("asistencia_req", ""))
    agregar_campo(doc, "7.3. % de Aprobación Exigido para Regularizar", datos_formulario.get("aprobacion_req", ""))
    agregar_campo(doc, "7.4. Cantidad de evaluaciones parciales", datos_formulario.get("cantidad_parciales", ""))
    agregar_campo(doc, "7.5. Modalidad y contenidos de evaluaciones sumativas durante la cursada", datos_formulario.get("eval_sumativas", ""))
    agregar_campo(doc, "7.6. Modalidad y contenidos de la evaluación complementaria para promoción", datos_formulario.get("eval_promocion", ""))
    agregar_campo(doc, "7.7. Criterio para decidir la calificación final por promoción", datos_formulario.get("calif_promocion", ""))

    # 8. ACREDITACIÓN POR EXAMEN FINAL
    agregar_titulo_seccion(doc, "8. ACREDITACIÓN POR EXAMEN FINAL")
    agregar_campo(doc, "8.1. Modalidad de examen final para Regulares", datos_formulario.get("final_reg_mod", ""))
    agregar_campo(doc, "8.2. Contenidos a evaluar en examen final Regular", datos_formulario.get("final_reg_cont", ""))
    agregar_campo(doc, "8.3. Criterio de calificación final para Regulares", datos_formulario.get("final_reg_crit", ""))
    agregar_campo(doc, "8.4. Modalidad de examen final para Libres", datos_formulario.get("final_lib_mod", ""))
    agregar_campo(doc, "8.5. Criterio de calificación final para Libres", datos_formulario.get("final_lib_crit", ""))

    # 9. BIBLIOGRAFÍA
    agregar_titulo_seccion(doc, "9. BIBLIOGRAFÍA")
    agregar_campo(doc, "9.1. Bibliografía Obligatoria", datos_formulario.get("bib_obligatoria", ""))
    agregar_campo(doc, "9.2. Bibliografía Complementaria", datos_formulario.get("bib_complementaria", ""))

    # 10. ADHESIÓN A PROGRAMAS ESPECIALES
    agregar_titulo_seccion(doc, "10. ADHESIÓN A PROGRAMAS ESPECIALES")
    progs = []
    if datos_formulario.get("prog_ltpa_cb"): progs.append("Condición LTPA (Libre con Trabajos Prácticos Aprobados)")
    if datos_formulario.get("prog_promuevo_cb"): progs.append('Programa "Yo También Promuevo"')
    if datos_formulario.get("prog_mesa_cb"): progs.append("Regularización en mesa de examen")
    if datos_formulario.get("prog_cuatrimestres_cb"): progs.append("Oferta de cursado en ambos cuatrimestres")
    if datos_formulario.get("prog_ninguno_cb"): progs.append("Ninguno de los anteriores")
    if datos_formulario.get("prog_otros_cb"): progs.append("Otros")

    agregar_campo(doc, "Programas Seleccionados", ", ".join(progs) if progs else "Ninguno")
    agregar_campo(doc, "Descripción de la Implementación", datos_formulario.get("desc_programas_text", ""))

    # Pie con fecha y hora de emisión
    doc.add_paragraph().paragraph_format.space_before = Pt(20)
    p_footer = doc.add_paragraph()
    p_footer.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    
    fecha_impresion = datetime.now().strftime("%d/%m/%Y a las %H:%M hs")
    run_foot = p_footer.add_run(f"Documento generado el {fecha_impresion} - Sistema PANDA (FBioyF - UNR)")
    run_foot.font.name = "Arial"
    run_foot.font.size = Pt(8)
    run_foot.font.italic = True
    run_foot.font.color.rgb = RGBColor(120, 120, 120)

    doc.save(ruta_salida)
    return ruta_salida

# Alias por compatibilidad
generar_doc_word = generar_docx_programa